import csv
import re
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook


CATEGORY_KEYWORDS = {
    "סופרמרקט": ["רמי לוי", "שופרסל", "מגה", "יינות", "סיטי", "סופר", "מכולת", "victory", "fresh", "מזון", "צרכנייה"],
    "דלק": ["דלק", "סונול", "פז", "תן", "fuel", "gas"],
    "מסעדות": ["מקדונלד", "בורגר", "פיצה", "שווארמה", "סושי", "מסעדה", "cafe", "קפה", "coffee", "wolt", "10bis", "אוכל בחוץ", "fat cow"],
    "בריאות": ["קופת חולים", "מכבי", "כללית", "ביטוח", "רופא", "תרופה", "pharmacy", "super-pharm", "רפואה", "קוסמטיקה"],
    "ביגוד": ["זארה", "H&M", "מנגו", "ASOS", "ביגוד", "נעליים", "next", "הלבשה"],
    "בידור": ["netflix", "spotify", "youtube", "סרט", "קולנוע", "theatre", "בידור", "פנאי"],
    "חינוך": ["בית ספר", "גן", "שכר לימוד", "ספרים", "קורס", "חינוך"],
    "תחבורה": ["אגד", "דן", "רב-קו", "מונית", "uber", "gett", "תחבורה", "רכב", "חניה"],
    "חשמל/מים": ["חשמל", "מים", "גז", "ועד בית", "ארנונה"],
    "תקשורת": ["סלקום", "פרטנר", "hot", "yes", "bezeq", "בזק", "תקשורת"],
    "קניות באינטרנט": ["aliexpress", "amazon", "ebay", "אלי אקספרס", "rue de"],
    "מוצרים לדירה": ["ikea", "איקאה", "רהיטים", "מוצרים לדירה"],
    "פנאי ובידור": ["פנאי", "בידור", "ספורט"],
    "עירייה וממשלה": ["עירייה", "ממשלה", "ארנונה", "מיסים"],
}


def guess_category(description: str) -> str:
    desc_lower = description.lower()
    for category, keywords in CATEGORY_KEYWORDS.items():
        for kw in keywords:
            if kw.lower() in desc_lower:
                return category
    return "כללי"


def parse_date(val) -> datetime | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    s = str(val).strip()
    for fmt in ["%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%m/%d/%Y"]:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def parse_amount(val) -> float | None:
    if val is None:
        return None
    try:
        cleaned = str(val).replace(",", "").replace("₪", "").replace(" ", "").strip()
        if not cleaned or cleaned in ["-", "—", ""]:
            return None
        result = float(cleaned)
        return abs(result) if result != 0 else None
    except (ValueError, TypeError):
        return None


def find_header_row(all_rows: list) -> tuple[int, dict]:
    """
    מחפש את שורת הכותרת ומחזיר (אינדקס, מיפוי עמודות).
    מחפש עמודות: תאריך, שם/תיאור, סכום, קטגוריה.
    """
    DATE_WORDS    = ["תאריך עסקה", "תאריך", "date"]
    DESC_WORDS    = ["שם בית העסק", "שם עסק", "תיאור", "פירוט", "description", "שם"]
    AMOUNT_WORDS  = ["סכום חיוב", "סכום", "amount", "חיוב", "זכות"]
    CAT_WORDS     = ["קטגוריה", "category"]

    for i, row in enumerate(all_rows[:30]):  # בדוק עד 30 שורות ראשונות
        cols = {}
        for j, cell in enumerate(row):
            if cell is None:
                continue
            cell_str = str(cell).strip()
            cell_low = cell_str.lower()

            if not cols.get("date") and any(w.lower() in cell_low for w in DATE_WORDS):
                cols["date"] = j
            if not cols.get("desc") and any(w.lower() in cell_low for w in DESC_WORDS):
                cols["desc"] = j
            if not cols.get("amount") and any(w.lower() in cell_low for w in AMOUNT_WORDS):
                cols["amount"] = j
            if not cols.get("cat") and any(w.lower() in cell_low for w in CAT_WORDS):
                cols["cat"] = j

        if cols.get("date") is not None and cols.get("amount") is not None:
            return i, cols

    return -1, {}


def parse_sheet_transactions(all_rows: list) -> list[dict]:
    """פרסור גיליון עסקאות (כרטיס אשראי וכדומה)"""
    header_idx, cols = find_header_row(all_rows)
    if header_idx == -1:
        return []

    rows_out = []
    for row in all_rows[header_idx + 1:]:
        if not row:
            continue

        # תאריך
        date_val = row[cols["date"]] if cols["date"] < len(row) else None
        date = parse_date(date_val)
        if not date:
            continue

        # סכום
        amount_val = row[cols["amount"]] if cols["amount"] < len(row) else None
        amount = parse_amount(amount_val)
        if not amount:
            continue

        # שם עסק
        desc_idx = cols.get("desc")
        description = str(row[desc_idx]).strip() if desc_idx is not None and desc_idx < len(row) and row[desc_idx] else "לא ידוע"
        if description in ["None", ""]:
            description = "לא ידוע"

        # קטגוריה
        cat_idx = cols.get("cat")
        if cat_idx is not None and cat_idx < len(row) and row[cat_idx]:
            category = str(row[cat_idx]).strip()
            if category in ["None", ""]:
                category = guess_category(description)
        else:
            category = guess_category(description)

        rows_out.append({
            "date": date.date(),
            "description": description,
            "amount": amount,
            "category": category,
            "person": "משותף",
        })

    return rows_out


def try_parse_financial_report(all_rows: list) -> list[dict]:
    """
    פרסור דוח פיננסי מסוג Pivot עם עמודות FY/LTM לכל שנה.
    """
    name_col = None
    name_row_idx = None

    for i, row in enumerate(all_rows[:20]):
        for j, cell in enumerate(row):
            if cell and "שם בית העסק" in str(cell):
                name_col = j
                name_row_idx = i
                break
        if name_col is not None:
            break

    if name_col is None:
        return []

    PERIODS = {
        "LTM-Feb26": 2026, "LTM Feb26": 2026, "LTM-Feb 26": 2026,
        "LTM-Sep23": 2023, "LTM Sep23": 2023,
        "FY25": 2025, "FY24": 2024, "FY23": 2023, "FY22": 2022,
    }

    period_cols = {}
    period_years = {}

    for i in range(max(0, name_row_idx - 4), min(len(all_rows), name_row_idx + 2)):
        for j, cell in enumerate(all_rows[i]):
            if not cell:
                continue
            cell_str = str(cell).strip()
            for period, year in PERIODS.items():
                if period.lower() in cell_str.lower() and period not in period_cols:
                    period_cols[period] = j
                    period_years[period] = year

    if not period_cols:
        return []

    preferred_order = [
        "LTM-Feb26", "LTM Feb26", "LTM-Feb 26",
        "FY25", "FY24", "FY23", "FY22",
        "LTM-Sep23", "LTM Sep23",
    ]

    rows_out = []
    for row in all_rows[name_row_idx + 1:]:
        if not row or len(row) <= name_col:
            continue
        name = row[name_col]
        if not name:
            continue
        name_str = str(name).strip()
        if not name_str or name_str == "None":
            continue
        if any(x in name_str for x in ['סה"כ', "סך הכל", "Total", "TOTAL", "סה''כ"]):
            continue

        for period in preferred_order:
            if period not in period_cols:
                continue
            col_j = period_cols[period]
            if col_j >= len(row):
                continue
            amount = parse_amount(row[col_j])
            if not amount:
                continue
            year = period_years[period]
            rows_out.append({
                "date": datetime(year, 6, 15).date(),
                "description": name_str,
                "amount": amount,
                "category": guess_category(name_str),
                "person": "משותף",
            })
            break

    return rows_out


def parse_excel(filepath: str) -> list[dict]:
    wb = load_workbook(filepath, data_only=True)
    rows = []

    # עדיפות לגיליונות עסקאות (לפי שם)
    TRANSACTION_SHEET_NAMES = [
        "כל הכרטיסים", "עסקאות", "transactions", "כרטיס",
        "רבעון", "2022", "2023", "2024", "2025",
    ]

    sheets = list(wb.worksheets)

    # מיון: גיליונות עסקאות קודם
    def sheet_priority(ws):
        name = ws.title.lower()
        for i, keyword in enumerate(TRANSACTION_SHEET_NAMES):
            if keyword.lower() in name:
                return i
        return 999

    sheets_sorted = sorted(sheets, key=sheet_priority)

    for ws in sheets_sorted:
        # דלג על גיליונות עזר
        skip_names = ["מקרא", "pivot", "ע\"ש", "עו\"ש", "סידור", "legend", "תקציב"]
        if any(s.lower() in ws.title.lower() for s in skip_names):
            continue

        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 3:
            continue

        # נסה פורמט עסקאות (כרטיס אשראי)
        tx_rows = parse_sheet_transactions(all_rows)
        if tx_rows:
            rows.extend(tx_rows)
            continue

        # נסה פורמט דוח פיננסי (Pivot)
        report_rows = try_parse_financial_report(all_rows)
        if report_rows:
            rows.extend(report_rows)

    return rows


def parse_csv(filepath: str) -> list[dict]:
    rows = []
    with open(filepath, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames or []

        col_date = col_desc = col_amount = None
        for h in headers:
            hl = h.lower()
            if any(w in hl for w in ["תאריך", "date"]):
                col_date = h
            elif any(w in hl for w in ["תיאור", "פירוט", "description", "שם"]):
                col_desc = h
            elif any(w in hl for w in ["סכום", "amount", "חיוב"]):
                if col_amount is None:
                    col_amount = h

        for row in reader:
            date = parse_date(row.get(col_date or "")) if col_date else datetime.today()
            if not date:
                continue
            raw_amount = row.get(col_amount, "0") if col_amount else list(row.values())[-1]
            amount = parse_amount(raw_amount)
            if amount is None:
                continue
            description = row.get(col_desc, "לא ידוע") if col_desc else list(row.values())[0]
            rows.append({
                "date": date.date() if hasattr(date, "date") else date,
                "description": str(description).strip(),
                "amount": amount,
                "category": guess_category(str(description)),
                "person": "משותף",
            })
    return rows


def parse_file(filepath: str) -> list[dict]:
    ext = Path(filepath).suffix.lower()
    if ext in [".xlsx", ".xls"]:
        return parse_excel(filepath)
    elif ext == ".csv":
        return parse_csv(filepath)
    elif ext == ".pdf":
        raise ValueError("ייבוא PDF אינו נתמך כרגע. אנא השתמש בקבצי Excel או CSV.")
    return []
